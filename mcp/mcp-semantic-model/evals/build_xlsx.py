"""Builds mail_model_graph_queries.xlsx (French) from the evaluated query set.

    python build_xlsx.py
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = Path(__file__).with_name("mail_model_graph_queries.xlsx")

HEADERS = ["Id", "Niveau", "Question", "Graphe requis", "Exploration", "Outils MCP", "Parcours dans le graphe",
           "Réponse attendue", "Résultat du test", "Vérifications"]

ROWS = [
    ("MG-01", "Simple", "Trouve le mail d'ouverture du dossier AV-2026-0135.", "Non", "Recherche d'un enregistrement (point d'entrée)",
     "find_records", "—",
     "Le mail « [AV-2026-0135] Outillage spécialisé indisponible … | Etape 01 - Position et preuves », envoyé par Theo Dubois le 06/01/2026.",
     "OK après correction : le mail est classé 1er (il était 2e)",
     "Le mail d'ouverture est classé avant les réponses « Re: Etape … » ; une référence seule ne renvoie que des enregistrements qui la portent."),
    ("MG-02", "Simple", "Qui a participé au mail d'ouverture du dossier AV-2026-0135 ?", "Oui", "Voisinage à 2 pas",
     "find_records, get_related_records", "Message e-mail → a pour participant → Participation → désigne → Correspondant",
     "3 correspondants : Karim Benali, Theo Dubois, Nora Dupont.", "OK",
     "Chemin Message → Participation → Correspondant ; nombre de liens exact = 3 ; aucun rôle affirmé (le rôle n'est pas renseigné dans les données)."),
    ("MG-03", "Simple", "Avec quoi le mail d'ouverture du dossier AV-2026-0135 est-il lié ?", "Oui", "Voisinage à 1 pas, toutes relations",
     "get_related_records", "Message e-mail → {Participation, Fait, Pièce jointe}",
     "3 participations, 1 fait (« Nous ouvrons AV-2026-0135 … », action, réalisé), 1 pièce jointe (AV-2026-0135-PJ-01.pdf).", "OK",
     "Les trois types de liens sont listés avec leurs nombres exacts ; aucun filtre de relation."),
    ("MG-04", "Simple", "Donne la pièce jointe du mail d'ouverture du dossier AV-2026-0135, avec sa taille.", "Oui", "Sélection par lien",
     "query_records (linked_to)", "Message e-mail → contient → Pièce jointe",
     "1 pièce jointe : AV-2026-0135-PJ-01.pdf, application/pdf, 51 590 octets.", "OK",
     "Total exact ; la taille est lue comme un nombre."),
    ("MG-05", "Simple", "Quels mails mentionnent des pénalités dans leur corps ? Combien ?", "Non", "Filtre sur un champ long",
     "query_records, find_records", "—", "300 mails dont le corps contient « pénalit ».", "OK",
     "Total exact = 300 ; find_records renvoie des passages du champ « corps »."),
    ("MG-06", "Simple", "Combien de mails par domaine d'expéditeur ?", "Non", "Regroupement et comptage",
     "query_records (group_by)", "—",
     "9 groupes : asteria-air 8 500, aerovia-maintenance 4 000, vectoris-components 2 500, pistea-handling 2 000, aeroflux-logistics 2 000, nautis-aero, reparo-aero, azurcover-aviation et altis-lease 1 500 chacun (total 25 000).",
     "OK", "La somme des groupes est égale au total ; comptes exacts, sans échantillonnage."),
    ("MG-07", "Moyen", "Évolution mensuelle du nombre de mails reçus en 2026.", "Non", "Regroupement par mois",
     "query_records (group_by, bucket month)", "—", "Janvier 4 793 ; février 4 857 ; mars 4 485 ; avril 3 888 ; mai 3 720 ; juin 3 257.",
     "OK", "6 tranches mensuelles ; la requête appliquée montre le regroupement par mois."),
    ("MG-08", "Moyen", "Combien de mails par thème de dossier, du plus fréquent au moins fréquent ?", "Non", "Comptages filtrés (un par thème)",
     "query_records", "—",
     "14 thèmes, total 25 000 : Pièce AOG livrée en retard 1 834 et Numéro de série incohérent 1 834 ; quatre thèmes à 1 833 ; huit thèmes à 1 750.",
     "OK (14 comptages filtrés ; le regroupement sur une partie d'un objet n'existe pas)",
     "Le total fait 25 000 ; les limites de l'outil sont signalées honnêtement."),
    ("MG-09", "Moyen", "Quels sont les faits au statut « contesté », par type de fait ?", "Non", "Filtre et regroupement",
     "query_records (group_by)", "—", "37 faits contestés répartis en 6 types (16, 10, 8, 1, 1, 1).", "OK",
     "Total = 37 ; la valeur du statut est prise dans la liste de valeurs du modèle."),
    ("MG-10", "Moyen", "Dans quels mails sont énoncés les faits contestés dont le texte mentionne un montant en EUR ?", "Oui", "Parcours à 1 pas depuis un filtre",
     "traverse_records", "Fait (contesté, contient « EUR ») → est énoncé dans → Message e-mail",
     "4 mails : AV-2026-0061, AV-2026-0145, AV-2026-0175 et AV-2026-0191.", "OK",
     "Le parcours utilise le nom inverse de la relation ; le texte du fait est lu dans le bon champ."),
    ("MG-11", "Moyen", "Quelles sont toutes les pénalités chiffrées énoncées dans les faits, avec le dossier concerné ?", "Non", "Filtre sur un champ texte",
     "query_records, find_records", "—",
     "3 faits : AV-2026-0210 « Pénalités SLA 156 200 EUR » (deux fois) et AV-2026-0074 « Pénalité retour 13 800, frais administratifs 3 450 (17 250 EUR) ».",
     "OK", "Les montants sont cités tels que stockés ; le doublon de AV-2026-0210 est restitué comme deux faits."),
    ("MG-12", "Moyen", "Quels dossiers ont une réserve de garantie énoncée par le courtier ?", "Oui", "Parcours à 1 pas depuis un filtre",
     "traverse_records", "Fait (« réserve la garantie ») → est énoncé dans → Message e-mail",
     "8 dossiers : AV-2026-0027, 0039, 0056, 0115, 0159, 0203, 0217 et 0244.", "OK (par parcours ; find_records seul renvoie 24 faits dont 16 sans rapport)",
     "Exactement 8, chacun relié à son mail par le lien (le numéro n'est pas deviné)."),
    ("MG-13", "Moyen", "Dans combien de dossiers Karim Benali apparaît-il comme participant, et lesquels ?", "Oui", "Parcours à 2 pas",
     "find_records, traverse_records", "Correspondant → a participé via → Participation → participe à → Message e-mail",
     "3 dossiers : AV-2026-0001, AV-2026-0135 et AV-2026-0267.", "OK",
     "Les noms de relations viennent de describe_model (libellés inverses) ; chemin Correspondant → Participation → Message."),
    ("MG-14", "Complexe", "Quels sont les autres correspondants qui ont partagé un dossier avec Karim Benali ?", "Oui", "Parcours à 4 pas (aller-retour)",
     "find_records, traverse_records", "Correspondant → Participation → Message → Participation → Correspondant",
     "7 correspondants : Amel Caron, Emma Petit, Lina Chevalier, Maxime Vidal, Nadia Caron, Nora Dupont, Theo Dubois ; Karim Benali lui-même exclu.",
     "OK", "Le retour sur le concept de départ exclut Karim Benali ; chaque co-participant est rattaché à un mail commun."),
    ("MG-15", "Complexe", "Quel lien y a-t-il entre Theo Dubois et Karim Benali ?", "Oui", "Plus court chemin entre deux enregistrements",
     "find_records, connect_records", "Theo Dubois → Participation → Message e-mail → Participation → Karim Benali",
     "Connectés en 4 liens, par le mail d'ouverture du dossier AV-2026-0135.", "OK",
     "Étapes réelles et orientées ; recherche complète."),
    ("MG-16", "Complexe", "Montre toutes les façons dont Karim Benali est relié à Nora Dupont, y compris via d'autres dossiers.", "Oui", "Chemins multiples entre deux enregistrements",
     "connect_records", "Karim Benali → … → Nora Dupont", "Connectés : 1 chemin de 4 liens (par AV-2026-0135).", "OK",
     "Chemins triés du plus court au plus long ; la complétude est signalée."),
    ("MG-17", "Complexe", "Reconstitue le fil du dossier AV-2026-0135 : les 10 premiers mails dans l'ordre chronologique, avec leur expéditeur.", "Non", "Filtre, tri et pagination",
     "query_records (order_by)", "—",
     "06/01 Theo Dubois ; 07/01 Amel Caron ; 07/01 Maxime Vidal ; 08/01 Inaya Martin ; 09/01 Alexandre Thomas ; 10/01 Camille Noel ; 11/01 Olivier Denis ; 12/01 Salma Renaud ; 13/01 Bastien Michel ; 15/01 Eva Henry (83 mails dans le fil).",
     "OK", "Ordre chronologique ; seul le mail d'ouverture a des liens dans le graphe, les réponses n'en ont pas."),
    ("MG-18", "Complexe", "Pour chaque domaine d'expéditeur, combien de mails d'ouverture (« Etape 01 ») ont une pièce jointe ?", "Oui", "Parcours à 1 pas avec regroupement sur le départ",
     "traverse_records (group_by étape 0)", "Message e-mail (ouverture) → contient → Pièce jointe",
     "300 mails d'ouverture : asteria-air 102, aerovia-maintenance 48, vectoris-components 30, puis 24, 24, 18, 18, 18, 18.", "OK",
     "Le total de départ est de 300 ; la somme des groupes est égale au total."),
    ("MG-19", "Complexe", "Les dossiers AV-2026-0061 et AV-2026-0191 ont chacun un fait contesté chiffré (22 400 EUR et 27 800 EUR). Sont-ils reliés entre eux ?", "Oui", "Parcours puis recherche de lien entre deux mails",
     "traverse_records, connect_records", "Fait → est énoncé dans → Message e-mail, puis Message ↔ Message",
     "Les deux faits sont retrouvés ; les deux mails ne sont pas connectés (recherche complète).", "OK",
     "« Non connectés » n'est affirmé que si la recherche est complète ; chaque affirmation renvoie à un enregistrement."),
    ("MG-20", "Complexe", "Parmi les dossiers à réserve de garantie, lesquels ont aussi un fait chiffré en EUR ?", "Oui", "Parcours à 2 pas (fait → mail → faits)",
     "traverse_records, get_related_records", "Fait (réserve de garantie) → est énoncé dans → Message e-mail → énonce → Fait (contient « EUR »)",
     "Aucun des 8 dossiers (intersection vide).", "OK",
     "Comptage à chaque étape : 8 mails, 0 avec un fait chiffré en EUR."),
]


def main() -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Requêtes"
    sheet.append(HEADERS)
    for row in ROWS:
        sheet.append(list(row))
    head = PatternFill("solid", fgColor="1F3A5F")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    yes, no = PatternFill("solid", fgColor="D9F2D9"), PatternFill("solid", fgColor="EDEDED")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        flag = row[3]
        flag.fill = yes if flag.value == "Oui" else no
        flag.font = Font(bold=True)
        flag.alignment = Alignment(horizontal="center", vertical="top")
    for index, width in enumerate([8, 11, 48, 12, 30, 28, 46, 60, 34, 56], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions

    summary = book.create_sheet("Synthèse")
    summary.append(["Niveau", "Requêtes", "Avec graphe", "Sans graphe"])
    for level in ("Simple", "Moyen", "Complexe"):
        rows = [r for r in ROWS if r[1] == level]
        summary.append([level, len(rows), sum(r[3] == "Oui" for r in rows), sum(r[3] == "Non" for r in rows)])
    summary.append(["Total", len(ROWS), sum(r[3] == "Oui" for r in ROWS), sum(r[3] == "Non" for r in ROWS)])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    for cell in summary[summary.max_row]:
        cell.font = Font(bold=True)
    for letter in "ABCD":
        summary.column_dimensions[letter].width = 16
    summary["A8"] = "Modèle : mail_model (version publiée) ; jeu de données : 25 000 mails de janvier à juin 2026, 300 dossiers AV-2026-NNNN."
    summary["A9"] = "« Graphe requis » = la réponse ne peut pas être lue dans les champs d'un seul enregistrement : elle exige de suivre des relations."
    book.save(OUT)


if __name__ == "__main__":
    main()
