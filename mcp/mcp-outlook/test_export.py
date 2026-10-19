"""Table-export checks — `python test_export.py`.

Only the pure parts: HTML-to-text, the CSV shape, and first-page extraction. The
Graph fetch and the Ceph write are covered through `export_emails_table` in
test_mail_enumeration.py with a patched transport.
"""
import io

import mail_export
from graph_helpers import logger


def test_html_to_text_strips_tags_and_unescapes():
    text = mail_export.html_to_text(
        "<div><p>Bonjour&nbsp;l'équipe</p><p>Contrat &amp; annexe</p></div>"
    )
    assert "<p>" not in text and "&amp;" not in text
    assert "Bonjour" in text and "Contrat & annexe" in text


def test_html_to_text_drops_scripts_and_styles():
    text = mail_export.html_to_text("<style>.x{color:red}</style><p>ok</p><script>evil()</script>")
    assert text.strip() == "ok"


def test_rows_use_aucune_when_there_is_no_attachment():
    row = mail_export.row_for(
        {"receivedDateTime": "2026-01-01T00:00:00Z", "subject": "S", "bodyPreview": "p",
         "body": {"contentType": "text", "content": "corps"}},
        [],
        "",
        "a@b.com",
    )
    assert row["Pièce jointe"] == "Aucune"
    assert row["Contenu de la première page de la pièce jointe"] == "Aucune"
    assert row["Contenu intégral"] == "corps"
    assert row["Expéditeur"] == "a@b.com"


def test_csv_has_a_bom_the_exact_columns_and_one_line_per_row():
    data = mail_export.build_csv([{"Objet": "Hello, world", "Expéditeur": "x@y.z"}])
    text = data.decode("utf-8-sig")
    header = text.splitlines()[0]
    assert header == ",".join(mail_export.COLUMNS)
    assert '"Hello, world"' in text  # comma inside the value is quoted


def test_markdown_table_keeps_the_columns_and_escapes_cells():
    md = mail_export.build_markdown([{"Objet": "a|b", "Contenu intégral": "l1\nl2", "Expéditeur": "x"}])
    lines = md.splitlines()
    assert lines[0] == "| " + " | ".join(mail_export.COLUMNS) + " |"
    assert lines[1].startswith("| ---")
    assert "a\\|b" in md
    assert "l1<br>l2" in md


def test_markdown_parts_split_a_long_table_but_rebuild_one():
    rows = [{"Objet": "x" * 200} for _ in range(300)]
    parts = mail_export.build_markdown_parts(rows, max_chars=4000)
    assert len(parts) > 1
    assert all(len(part) <= 4200 for part in parts)
    header = "| " + " | ".join(mail_export.COLUMNS) + " |"
    assert parts[0].splitlines()[0] == header
    assert "---" not in parts[1]  # only the first part carries the header row
    data_lines = sum(1 for part in parts for line in part.splitlines() if "---" not in line)
    assert data_lines == 301  # header + 300 rows


def test_unknown_attachment_type_is_noted_not_guessed():
    out = mail_export.first_page_of(b"binary", "image/png", "logo.png")
    assert "non extrait" in out


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _has_libreoffice():
    return mail_export._libreoffice_binary() is not None


def _docx_bytes(first: str, second: str = "", page_break: bool = False) -> bytes:
    import docx
    from docx.enum.text import WD_BREAK

    document = docx.Document()
    document.add_paragraph(first)
    if second:
        if page_break:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        document.add_paragraph(second)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_render_reports_unavailable_without_libreoffice():
    original = mail_export._libreoffice_binary
    mail_export._libreoffice_binary = lambda: None
    try:
        assert mail_export.render_docx_to_pdf(b"anything") is None
        out = mail_export.first_page_of(b"anything", DOCX_MIME, "x.docx")
    finally:
        mail_export._libreoffice_binary = original
    assert out == mail_export.DOCX_RENDER_UNAVAILABLE


def test_docx_first_page_reads_the_leading_text():
    if not _has_libreoffice():
        logger.info("    (skip: LibreOffice not installed)")
        return
    out = mail_export.first_page_of(
        _docx_bytes("Contrat de prestation", "Article 1 — objet"), DOCX_MIME, "contrat.docx"
    )
    assert "Contrat de prestation" in out and "Article 1" in out


def test_docx_first_page_includes_table_text():
    if not _has_libreoffice():
        logger.info("    (skip: LibreOffice not installed)")
        return
    import docx

    document = docx.Document()
    document.add_paragraph("Intro")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Clause de notification"
    table.rows[0].cells[1].text = "30 jours"
    buffer = io.BytesIO()
    document.save(buffer)
    out = mail_export.first_page_of(buffer.getvalue(), DOCX_MIME, "c.docx")
    assert "Clause de notification" in out and "30 jours" in out


def test_docx_first_page_excludes_the_second_page():
    if not _has_libreoffice():
        logger.info("    (skip: LibreOffice not installed)")
        return
    data = _docx_bytes("PREMIERE PAGE MARQUEUR", "DEUXIEME PAGE MARQUEUR", page_break=True)
    out = mail_export.first_page_of(data, DOCX_MIME, "deux-pages.docx")
    assert "PREMIERE PAGE MARQUEUR" in out
    assert "DEUXIEME PAGE MARQUEUR" not in out


def test_safe_filename_keeps_a_usable_stem():
    assert mail_export.safe_filename("emails_ssnousi@yellowsys.fr") == "emails_ssnousi_yellowsys.fr"
    assert mail_export.safe_filename("") == "emails"
    assert "/" not in mail_export.safe_filename("a/b c")


if __name__ == "__main__":
    passed = 0
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            passed += 1
            logger.info(f"  ok  {name}")
    logger.info(f"\n{passed} passed")
