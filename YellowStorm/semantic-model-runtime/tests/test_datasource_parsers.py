"""Bounded tabular preview parsers (P3.7-P3.10).

Synthetic bytes only: no network, no storage, no NestJS. XLSX fixtures are
generated in-test with openpyxl (write path); the parser under test only ever
uses the lazy read path.
"""

from __future__ import annotations

import io

import pytest

from app.datasource import parsers
from app.datasource.discovery import preview_source
from app.datasource.parsers import (
    PREVIEW_MAX_COLUMNS,
    PREVIEW_SAMPLE_ROWS,
    SCAN_ROW_LIMIT,
    SHEET_ROW_KEY,
    parse_csv_preview,
    parse_xlsx_preview,
)


def _csv(rows: list[list[str]], delimiter: str = ",") -> bytes:
    out = io.StringIO()
    out.write(delimiter.join([]) if False else "")
    import csv as _csvmod

    writer = _csvmod.writer(out, delimiter=delimiter, lineterminator="\n")
    writer.writerows(rows)
    return out.getvalue().encode("utf-8")


def _xlsx(sheets: dict[str, list[list]]) -> bytes:
    openpyxl = pytest.importorskip("openpyxl")
    workbook = openpyxl.Workbook()
    first = True
    for name, rows in sheets.items():
        sheet = workbook.active if first else workbook.create_sheet(name)
        first = False
        sheet.title = name
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_csv_basic_preserves_leading_zeroes_and_suggests_types():
    data = _csv([["customer_id", "name", "balance"],
                 ["00123", "Acme", "10.5"],
                 ["00456", "Globex", "20"]])
    parsed = parse_csv_preview(data)
    assert parsed["structure"]["kind"] == "csv"
    assert parsed["structure"]["delimiter"] == ","
    assert parsed["structure"]["headerRow"] == 1
    assert parsed["structure"]["dataRows"] == 2
    assert parsed["coverage"] == {"sampled": True, "completeProfileDone": True}
    assert parsed["samples"][0] == {"customer_id": "00123", "name": "Acme",
                                    "balance": "10.5", SHEET_ROW_KEY: 2}
    assert parsed["samples"][1][SHEET_ROW_KEY] == 3
    profiles = {p["name"]: p for p in parsed["fieldProfiles"]}
    assert profiles["customer_id"]["suggestedType"] == "number"
    assert profiles["name"]["type"] == "text"
    assert any(w["code"] == "header_row_assumed" for w in parsed["warnings"])
    assert parsed["contentFingerprint"].startswith("sha256:")


def test_csv_bom_semicolon_and_multiline_records():
    raw = "réf;libellé\n\"a;b\";\"line1\nline2\"\n".encode("utf-8-sig")
    parsed = parse_csv_preview(raw)
    assert parsed["structure"]["delimiter"] == ";"
    assert parsed["structure"]["encoding"] == "utf-8-sig"
    assert parsed["structure"]["bomPresent"] is True
    assert any(w["code"] == "delimiter_sniffed" for w in parsed["warnings"])
    assert len(parsed["samples"]) == 1
    assert parsed["samples"][0]["libellé"] == "line1\nline2"
    # Source row number tracks the physical line where the record ends.
    assert parsed["samples"][0][SHEET_ROW_KEY] == 3


def test_csv_quoted_delimiters_do_not_vote():
    # The commas live inside quotes; the file is semicolon-delimited.
    raw = 'a;b\nx;"p,q,r"\n'.encode("utf-8")
    parsed = parse_csv_preview(raw)
    assert parsed["structure"]["delimiter"] == ";"
    assert parsed["samples"][0] == {"a": "x", "b": "p,q,r", SHEET_ROW_KEY: 2}


def test_csv_strict_mode_rejects_unterminated_quotes():
    with pytest.raises(ValueError, match="csv_malformed"):
        parse_csv_preview(b'a,b\n"x,y\n')


def test_csv_header_only_file_infers_delimiter():
    parsed = parse_csv_preview(b"a;b\n")
    assert parsed["structure"]["delimiter"] == ";"
    assert parsed["structure"]["columns"] == ["a", "b"]
    assert parsed["structure"]["dataRows"] == 0
    assert parsed["coverage"] == {"sampled": True, "completeProfileDone": True}


def test_csv_newline_dense_prefix_is_bounded():
    import time

    data = (b"\n" * 200_000) + b"a,b\n1,2\n"
    started = time.perf_counter()
    with pytest.raises(ValueError, match="csv_no_header_row"):
        parse_csv_preview(data)
    assert time.perf_counter() - started < 5

    # Ordinary leading blanks still work.
    parsed = parse_csv_preview(b"\n\na,b\n1,2\n")
    assert parsed["structure"]["columns"] == ["a", "b"]
    assert parsed["samples"][0][SHEET_ROW_KEY] == 4


def test_csv_multiline_record_trips_the_physical_bound(monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(parsers, "SCAN_ROW_LIMIT", 10)
    # One quoted record spans 50 physical lines, then valid rows follow.
    body = "note\n" + '"line1\n' + "\n".join(f"line{i}" for i in range(2, 51)) + '"\n'
    tail = "".join(f"tail{i}\n" for i in range(5))
    parsed = parse_csv_preview((body + tail).encode("utf-8"))
    assert parsed["coverage"] == {"sampled": True, "completeProfileDone": False}
    assert any(w["code"] == "scan_capped" for w in parsed["warnings"])
    assert parsed["scannedRows"] == 1
    # Traversal stopped at the bound: rows past it are absent.
    assert len(parsed["samples"]) == 1
    assert "line50" in parsed["samples"][0]["note"]


def test_csv_record_starting_past_the_cap_is_never_consumed(monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(parsers, "SCAN_ROW_LIMIT", 10)
    head = "id\n" + "".join(f"{i}\n" for i in range(10))
    # Raw bytes, NOT csv.writer: the writer would escape these into valid
    # records and the test could not fail on the pre-fix implementation.
    # A valid multiline record starting past the filled cap is never fetched.
    parsed = parse_csv_preview((head + '"multi\nline"\n' + "tail\n").encode("utf-8"))
    assert parsed["coverage"] == {"sampled": True, "completeProfileDone": False}
    assert parsed["scannedRows"] == 10
    assert len(parsed["samples"]) == 10
    assert all("multi" not in str(row) for row in parsed["samples"])
    # A malformed record starting past the cap is never consumed either: the
    # old implementation fetched it and raised instead of returning partial.
    parsed = parse_csv_preview((head + '"unterminated\n').encode("utf-8"))
    assert parsed["coverage"] == {"sampled": True, "completeProfileDone": False}
    assert parsed["scannedRows"] == 10


def test_csv_physical_bound_is_inclusive(monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(parsers, "SCAN_ROW_LIMIT", 10)
    # Fewer than 10 data rows, but enough blank physical lines to reach the
    # exact physical bound, then a raw malformed record. The inclusive bound
    # stops at line SCAN+1 without fetching what follows.
    body = "id\n" + "1\n2\n3\n" + "\n" * 7 + '"unterminated\n'
    parsed = parse_csv_preview(body.encode("utf-8"))
    assert parsed["coverage"] == {"sampled": True, "completeProfileDone": False}
    assert parsed["scannedRows"] == 3
    assert [row["id"] for row in parsed["samples"]] == ["1", "2", "3"]


def test_csv_provenance_key_is_reserved():
    parsed = parse_csv_preview(_csv([["__sheetRow", "id"], ["evil", "1"]]))
    assert parsed["structure"]["columns"] == ["sheetRow", "id"]
    row = parsed["samples"][0]
    assert row[SHEET_ROW_KEY] == 2
    assert row["sheetRow"] == "evil"
    assert row["id"] == "1"
    assert {p["name"] for p in parsed["fieldProfiles"]} == {"sheetRow", "id"}


def test_xlsx_provenance_key_is_reserved():
    parsed = parse_xlsx_preview(_xlsx({"S": [["__sheetRow", "id"], ["evil", "1"]]}))
    assert parsed["structure"]["columns"] == ["sheetRow", "id"]
    row = parsed["samples"][0]
    assert row[SHEET_ROW_KEY] == 2
    assert row["sheetRow"] == "evil"
    assert {p["name"] for p in parsed["fieldProfiles"]} == {"sheetRow", "id"}
    parsed = parse_csv_preview(_csv([["id", "id", "id__2"], ["1", "2", "3"]]))
    assert parsed["structure"]["columns"] == ["id", "id__2", "id__2__2"]
    assert any(w["code"] == "headers_renamed" for w in parsed["warnings"])
    assert parsed["samples"][0] == {"id": "1", "id__2": "2", "id__2__2": "3",
                                    SHEET_ROW_KEY: 2}
    empty_first = parse_csv_preview(_csv([["", "Column 1"], ["a", "b"]]))
    assert empty_first["structure"]["columns"] == ["Column 1", "Column 1__2"]


def test_csv_large_valid_field_parses_and_truncates_in_samples():
    big = "x" * 200_000  # above the stdlib 128 KiB default, below our 10 MB bound
    parsed = parse_csv_preview(_csv([["note"], [big]]))
    assert parsed["samples"][0]["note"].endswith("…")
    assert len(parsed["samples"][0]["note"]) == 501
    assert any(w["code"] == "sample_cells_truncated" for w in parsed["warnings"])


def test_xlsx_sparse_huge_range_is_bounded_and_partial():
    openpyxl = pytest.importorskip("openpyxl")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sparse"
    sheet.append(["id"])
    sheet.append(["top"])
    sheet.cell(row=6000, column=1, value="far")
    sheet.cell(row=2, column=200, value="wide")
    buffer = io.BytesIO()
    workbook.save(buffer)
    parsed = parse_xlsx_preview(buffer.getvalue())
    # Stops at the physical traversal bound instead of walking the full range.
    assert parsed["coverage"] == {"sampled": True, "completeProfileDone": False}
    assert any(w["code"] == "scan_capped" for w in parsed["warnings"])
    # Width comes from populated header cells: no synthetic padding columns.
    # Values past the header width are outside the previewed schema and are
    # omitted, matching the existing mapping behavior.
    assert parsed["structure"]["columns"] == ["id"]
    assert parsed["samples"][0]["id"] == "top"


def test_csv_encoding_fallback_and_size_gate():
    parsed = parse_csv_preview("café;1\n".encode("latin-1"))
    assert parsed["structure"]["encoding"] == "latin-1"
    assert any(w["code"] == "encoding_fallback_latin1" for w in parsed["warnings"])
    with pytest.raises(ValueError, match="source_too_large"):
        parse_csv_preview(b"x" * (50 * 1024 * 1024 + 1))
    with pytest.raises(ValueError, match="csv_no_header_row"):
        parse_csv_preview(b"\n\n")


def test_csv_bounds_samples_columns_and_scan_cap(monkeypatch: pytest.MonkeyPatch):
    wide = [["c%d" % i for i in range(PREVIEW_MAX_COLUMNS + 5)],
            ["v%d" % i for i in range(PREVIEW_MAX_COLUMNS + 5)]]
    parsed = parse_csv_preview(_csv(wide))
    assert len(parsed["structure"]["columns"]) == PREVIEW_MAX_COLUMNS
    assert any(w["code"] == "columns_truncated" for w in parsed["warnings"])

    rows = [["id"]] + [[str(i)] for i in range(PREVIEW_SAMPLE_ROWS + 10)]
    parsed = parse_csv_preview(_csv(rows))
    assert len(parsed["samples"]) == PREVIEW_SAMPLE_ROWS

    monkeypatch.setattr(parsers, "SCAN_ROW_LIMIT", 10)
    capped = parse_csv_preview(_csv([["id"]] + [[str(i)] for i in range(50)]))
    assert capped["coverage"] == {"sampled": True, "completeProfileDone": False}
    assert capped["structure"]["dataRows"] is None
    assert any(w["code"] == "scan_capped" for w in capped["warnings"])
    assert SCAN_ROW_LIMIT == 5000  # the real cap is untouched elsewhere


def test_xlsx_sheets_selection_and_types():
    from datetime import datetime

    data = _xlsx({"Customers": [["id", "name", "since"],
                                ["001", "Acme", datetime(2026, 1, 2, 3, 4, 5)],
                                ["002", "Globex", datetime(2026, 5, 6, 0, 0, 0)]],
                  "Hidden": [["a"], ["1"]]})
    parsed = parse_xlsx_preview(data)
    assert parsed["structure"]["kind"] == "xlsx"
    assert [s["name"] for s in parsed["structure"]["sheets"]] == ["Customers", "Hidden"]
    assert parsed["structure"]["selectedSheet"] == "Customers"
    assert parsed["structure"]["dataRows"] == 2
    assert parsed["samples"][0]["id"] == "001"
    assert parsed["samples"][0]["since"] == "2026-01-02T03:04:05"
    profiles = {p["name"]: p for p in parsed["fieldProfiles"]}
    assert profiles["since"]["type"] == "date"
    assert any(w["code"] == "header_row_assumed" for w in parsed["warnings"])

    second = parse_xlsx_preview(data, {"sheetName": "Hidden"})
    assert second["structure"]["selectedSheet"] == "Hidden"
    with pytest.raises(ValueError, match="sheet_not_found"):
        parse_xlsx_preview(data, {"sheetName": "Missing"})
    with pytest.raises(ValueError, match="unreadable_source"):
        parse_xlsx_preview(b"not a workbook")


def test_xlsx_iterator_not_advanced_past_budget(monkeypatch: pytest.MonkeyPatch):
    pytest.importorskip("openpyxl")
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet

    data_yields = {"n": 0}
    real_iter_rows = ReadOnlyWorksheet.iter_rows

    def counting_iter_rows(self, *args, **kwargs):
        if kwargs.get("min_row", 1) == 2 and kwargs.get("values_only", False):
            for row in real_iter_rows(self, *args, **kwargs):
                data_yields["n"] += 1
                yield row
        else:
            yield from real_iter_rows(self, *args, **kwargs)

    monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", counting_iter_rows)
    monkeypatch.setattr(parsers, "SCAN_ROW_LIMIT", 10)
    parsed = parse_xlsx_preview(_xlsx({"S": [["id"]] + [[str(i)] for i in range(15)]}))
    assert parsed["scannedRows"] == 10
    assert parsed["coverage"] == {"sampled": True, "completeProfileDone": False}
    # The data pass fetched exactly the budget: the 11th row was never parsed.
    assert data_yields["n"] == 10


def test_xlsx_sparse_iterator_not_advanced_past_budget(monkeypatch: pytest.MonkeyPatch):
    pytest.importorskip("openpyxl")
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet

    data_yields = {"n": 0}
    real_iter_rows = ReadOnlyWorksheet.iter_rows

    def counting_iter_rows(self, *args, **kwargs):
        if kwargs.get("min_row", 1) == 2 and kwargs.get("values_only", False):
            for row in real_iter_rows(self, *args, **kwargs):
                data_yields["n"] += 1
                yield row
        else:
            yield from real_iter_rows(self, *args, **kwargs)

    monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", counting_iter_rows)
    monkeypatch.setattr(parsers, "SCAN_ROW_LIMIT", 10)
    openpyxl = pytest.importorskip("openpyxl")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sparse"
    sheet.append(["id"])
    for value in ("a", "b", "c"):
        sheet.append([value])
    sheet.cell(row=100, column=1, value="far")
    buffer = io.BytesIO()
    workbook.save(buffer)
    parsed = parse_xlsx_preview(buffer.getvalue())
    assert parsed["scannedRows"] == 3
    assert parsed["coverage"] == {"sampled": True, "completeProfileDone": False}
    # Sparse traversal stops at the physical budget, not the reported range.
    assert data_yields["n"] == 10


def test_xlsx_narrow_and_blank_sheets_have_no_synthetic_columns():
    narrow = parse_xlsx_preview(_xlsx({"One": [["only"], ["1"], ["2"]]}))
    assert narrow["structure"]["columns"] == ["only"]
    assert narrow["structure"]["dataRows"] == 2
    assert narrow["coverage"] == {"sampled": True, "completeProfileDone": True}
    assert len(narrow["samples"][0]) == 2  # value + __sheetRow, nothing padded

    blank = parse_xlsx_preview(_xlsx({"Empty": []}))
    assert blank["structure"]["columns"] == []
    assert blank["samples"] == []
    assert any(w["code"] == "empty_sheet" for w in blank["warnings"])


def test_byte_path_skips_terminal_discovery_states():
    base = {"workspaceId": "6512f0a1c9e77a001234aaa1",
            "assetId": "6512f0a1c9e77a001234bbb2",
            "originalName": "customer-registry.xlsx",
            "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "sizeBytes": 184320,
            "contentHash": "9f2a1c0e5b6d7a8b9c0d1e2f3a4b5c6d",
            "uploadedAt": "2026-09-10T08:30:00.000Z",
            "indexingStatus": None}
    data = _xlsx({"Customers": [["customer_id"], ["00123"]]})
    for override, status in [({"protected": True}, "protected"),
                             ({"corrupt": True}, "corrupt"),
                             ({"mimeType": "application/x-msaccess"}, "unsupported")]:
        result = preview_source({**base, **override}, None, data)
        assert result["profile"]["status"] == status
        assert result["profile"]["samples"] == []
        assert "fieldProfiles" not in result

    doc = {**base, "mimeType": "application/pdf", "indexingStatus": "failed"}
    result = preview_source(doc, None, data)
    assert result["profile"]["status"] == "indexing_required"
    assert result["profile"]["samples"] == []

    # Oversized metadata keeps its partial status and warning even when the
    # supplied bytes are smaller: parsing must never upgrade it to ready.
    huge = {**base, "sizeBytes": 60 * 1024 * 1024}
    result = preview_source(huge, None, data)
    assert result["profile"]["status"] == "partial"
    assert result["profile"]["samples"] == []
    assert any(w["code"] == "source_too_large_for_preview"
               for w in result["profile"]["warnings"])
    assert "fieldProfiles" not in result
    openpyxl = pytest.importorskip("openpyxl")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Calc"
    sheet.append(["a", "b", "total"])
    sheet.append([1, 2, "=A2+B2"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    parsed = parse_xlsx_preview(buffer.getvalue())
    assert any(w["code"] == "formula_values_are_cached" for w in parsed["warnings"])
    assert any(w["code"] == "no_formula_execution" for w in parsed["warnings"])


def test_preview_source_byte_path_merges_into_schema_valid_profile():
    from pathlib import Path
    import json

    jsonschema = pytest.importorskip("jsonschema")
    from referencing import Registry, Resource

    source = {"workspaceId": "6512f0a1c9e77a001234aaa1",
              "assetId": "6512f0a1c9e77a001234bbb2",
              "originalName": "customer-registry.xlsx",
              "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
              "sizeBytes": 184320,
              "contentHash": "9f2a1c0e5b6d7a8b9c0d1e2f3a4b5c6d",
              "uploadedAt": "2026-09-10T08:30:00.000Z",
              "indexingStatus": None}
    data = _xlsx({"Customers": [["customer_id", "name"], ["00123", "Acme"]]})
    result = preview_source(source, None, data)
    profile = result["profile"]
    assert profile["status"] == "ready"
    assert profile["coverage"] == {"sampled": True, "completeProfileDone": True}
    assert profile["samples"][0]["customer_id"] == "00123"
    assert result["ingestionPlan"]["decision"] == "prepare_dataset"
    assert result["contentFingerprint"].startswith("sha256:")
    assert result["fieldProfiles"]

    contracts = Path(__file__).resolve().parents[1] / "contracts" / "v1"
    asset = json.loads((contracts / "asset-ref.schema.json").read_text(encoding="utf-8"))
    schema = json.loads((contracts / "discovery-profile.schema.json").read_text(encoding="utf-8"))
    registry = Registry().with_resources([
        (asset["$id"], Resource.from_contents(asset)),
        (schema["$id"], Resource.from_contents(schema))])
    jsonschema.Draft7Validator(schema, registry=registry).validate(profile)

    # Bytes for a non-tabular source are not parsed: the terminal metadata
    # profile is returned unchanged instead of overriding its status.
    doc = {**source, "mimeType": "application/pdf"}
    skipped = preview_source(doc, None, data)
    assert skipped["profile"]["status"] == "indexing_required"
    assert skipped["profile"]["samples"] == []
    assert "fieldProfiles" not in skipped
    # Without bytes the metadata-only path is unchanged.
    assert preview_source(source)["profile"]["samples"] == []


def test_parsers_module_import_pulls_no_heavy_libs():
    """Hermetic: importing the parsers module must not initialize openpyxl,
    whatever sibling tests already imported in this process."""
    import subprocess
    import sys as _sys
    from pathlib import Path

    code = ("import sys, app.datasource.parsers; "
            "assert 'openpyxl' not in sys.modules, 'openpyxl'; "
            "assert 'duckdb' not in sys.modules, 'duckdb'; "
            "print('lazy-ok')")
    completed = subprocess.run(
        [_sys.executable, "-c", code],
        cwd=Path(__file__).resolve().parents[1],
        capture_output=True, text=True, timeout=60,
    )
    assert completed.returncode == 0, completed.stderr
    assert "lazy-ok" in completed.stdout
