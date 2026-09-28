"""Full tabular preparation into versioned Parquet datasets."""

from __future__ import annotations

import csv
import hashlib
import io
from pathlib import Path
from typing import Any, Iterable

from .discovery import discover
from .parsers import (SHEET_ROW_KEY, Stringify,
                      _decode_text, _disambiguate_headers, _sniff_delimiter,
                      parse_xlsx_preview)

BATCH_ROWS = 1000
MAX_DATASET_BYTES = 200 * 1024 * 1024


def dataset_id(source: dict[str, Any], options: dict[str, Any] | None) -> str:
    profile = discover(source, options)
    identity = f"{profile['assetRef']['assetVersionId']}|{profile['parserFingerprint']}"
    return f"ds_{hashlib.sha256(identity.encode()).hexdigest()[:24]}"


def _write_batches(headers: list[str], rows: Iterable[tuple[int, list[Any]]], output: Path) -> int:
    import pyarrow as pa
    import pyarrow.parquet as pq

    schema = pa.schema([(SHEET_ROW_KEY, pa.int64()), *[(name, pa.string()) for name in headers]])
    writer = pq.ParquetWriter(output, schema, compression="zstd")
    count = 0
    batch: list[tuple[int, list[Any]]] = []
    try:
        for row in rows:
            batch.append(row)
            if len(batch) >= BATCH_ROWS:
                count += _write_batch(writer, schema, headers, batch)
                batch.clear()
        if batch:
            count += _write_batch(writer, schema, headers, batch)
    finally:
        writer.close()
    return count


def _write_batch(writer: Any, schema: Any, headers: list[str], rows: list[tuple[int, list[Any]]]) -> int:
    import pyarrow as pa

    columns: dict[str, list[Any]] = {SHEET_ROW_KEY: [number for number, _ in rows]}
    for index, name in enumerate(headers):
        columns[name] = [None if index >= len(values) or values[index] is None
                         else Stringify(values[index]) for _, values in rows]
    writer.write_table(pa.Table.from_pydict(columns, schema=schema))
    return len(rows)


def _csv_rows(data: bytes, options: dict[str, Any]) -> tuple[list[str], Iterable[tuple[int, list[Any]]]]:
    warnings: list[dict[str, str]] = []
    text, _ = _decode_text(data, options, warnings)
    delimiter = _sniff_delimiter(text, options, warnings)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter, strict=True)
    raw_headers = next((row for row in reader if any(cell.strip() for cell in row)), [])
    headers = _disambiguate_headers(raw_headers, warnings)

    def rows() -> Iterable[tuple[int, list[Any]]]:
        try:
            for values in reader:
                if any(cell != "" for cell in values):
                    yield reader.line_num, values[:len(headers)]
        except csv.Error as exc:
            raise ValueError("csv_malformed") from exc

    return headers, rows()


def _xlsx_rows(data: bytes, options: dict[str, Any]) -> tuple[list[str], Iterable[tuple[int, list[Any]]], Any]:
    import openpyxl

    preview = parse_xlsx_preview(data, options)
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook[preview["structure"]["selectedSheet"]]
    # Workbooks written without a <dimension> report no max_column; read the header row to its end then.
    header_cells = list(next(sheet.iter_rows(min_row=1, max_row=1,
                                             max_col=max(sheet.max_column or 0, 1) if sheet.max_column else None,
                                             values_only=True), []))
    while header_cells and header_cells[-1] is None:
        header_cells.pop()
    warnings: list[dict[str, str]] = []
    headers = _disambiguate_headers(
        ["" if value is None else Stringify(value) for value in header_cells], warnings)

    def rows() -> Iterable[tuple[int, list[Any]]]:
        for number, values in enumerate(
                sheet.iter_rows(min_row=2, max_col=max(len(headers), 1), values_only=True), start=2):
            cells = list(values)
            if any(value is not None and Stringify(value).strip() for value in cells):
                yield number, cells[:len(headers)]

    return headers, rows(), workbook


def prepare_parquet(source: dict[str, Any], options: dict[str, Any] | None,
                    data: bytes, output: Path) -> dict[str, Any]:
    opts = options or {}
    mime = source.get("mimeType")
    workbook = None
    if mime == "text/csv":
        headers, rows = _csv_rows(data, opts)
    elif mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        headers, rows, workbook = _xlsx_rows(data, opts)
    else:
        raise ValueError("unsupported_format_for_dataset")
    try:
        row_count = _write_batches(headers, rows, output)
    finally:
        if workbook is not None:
            workbook.close()
    size = output.stat().st_size
    if size > MAX_DATASET_BYTES:
        output.unlink(missing_ok=True)
        raise ValueError("dataset_too_large")
    digest = hashlib.sha256()
    with output.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return {"datasetId": dataset_id(source, options), "rowCount": row_count,
            "columns": [SHEET_ROW_KEY, *headers], "sizeBytes": size,
            "contentHash": f"sha256:{digest.hexdigest()}"}
