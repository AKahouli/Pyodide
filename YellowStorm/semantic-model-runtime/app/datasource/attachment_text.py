"""Plain text of an e-mail attachment, so its content becomes a field of the record.

Reads the formats the Email Reader marks as readable (PDF, Word, Excel, PowerPoint, CSV, text). A file that
yields no text (a scanned PDF, a damaged file) gives ``None``: the record keeps its other
fields and says why in ``statut_contenu``.
"""

from __future__ import annotations

import io
import re
import zipfile
from xml.etree import ElementTree

from .extraction_limits import CONTENT_CHARS, ExtractionLimits

_WORD = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_DRAWING = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


class InflatedTooLarge(Exception):
    """A zipped office file that holds more than the admin's limit once decompressed."""


def _inflated(archive: zipfile.ZipFile, names: list[str], limits: ExtractionLimits) -> None:
    # The header size is what the reader is allowed to read back, so it bounds what memory receives.
    if sum(archive.getinfo(name).file_size for name in names) > limits.inflated_bytes:
        raise InflatedTooLarge


def _pdf(payload: bytes, limits: ExtractionLimits) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(payload))
    pages = []
    for page in reader.pages[:limits.pdf_pages]:
        pages.append(page.extract_text() or "")
    return "\n\n".join(pages)


def _docx(payload: bytes, limits: ExtractionLimits) -> str:
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        _inflated(archive, ["word/document.xml"], limits)
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    paragraphs = []
    for paragraph in root.iter(f"{_WORD}p"):
        parts = []
        for node in paragraph.iter():
            if node.tag == f"{_WORD}t" and node.text:
                parts.append(node.text)
            elif node.tag == f"{_WORD}tab":
                parts.append("\t")
            elif node.tag in (f"{_WORD}br", f"{_WORD}cr"):
                parts.append("\n")
        paragraphs.append("".join(parts))
    return "\n".join(paragraphs)


def _xlsx(payload: bytes, limits: ExtractionLimits) -> str:
    import openpyxl

    # openpyxl streams the sheets (bounded by the cell limit) but loads the shared strings, styles and
    # workbook parts whole: those are what a crafted file can inflate.
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        whole = {"xl/sharedStrings.xml", "xl/styles.xml", "xl/workbook.xml"}
        _inflated(archive, [name for name in archive.namelist() if name in whole], limits)
    workbook = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        lines, cells = [], 0
        for sheet in workbook.worksheets:
            lines.append(f"## {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                while values and not values[-1]:
                    values.pop()
                if values:
                    lines.append("\t".join(values))
                cells += len(values)
                if cells >= limits.sheet_cells:
                    return "\n".join(lines)
        return "\n".join(lines)
    finally:
        workbook.close()


def _slide_number(name: str) -> int:
    return int(re.search(r"(\d+)\.xml$", name).group(1))  # type: ignore[union-attr]


def _paragraphs(xml: bytes) -> list[str]:
    lines = []
    for paragraph in ElementTree.fromstring(xml).iter(f"{_DRAWING}p"):
        text = "".join(node.text for node in paragraph.iter(f"{_DRAWING}t") if node.text)
        if text.strip():
            lines.append(text)
    return lines


def _pptx(payload: bytes, limits: ExtractionLimits) -> str:
    """Slide texts in order, each under a ``## Slide n`` heading, followed by that slide's speaker notes."""
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        names = set(archive.namelist())
        slides = sorted((name for name in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)),
                        key=_slide_number)[:limits.slides]
        notes_of = {name: f"ppt/notesSlides/notesSlide{_slide_number(name)}.xml" for name in slides}
        _inflated(archive, [*slides, *(notes for notes in notes_of.values() if notes in names)], limits)
        sections = []
        for name in slides:
            lines = [f"## Slide {_slide_number(name)}", *_paragraphs(archive.read(name))]
            if notes_of[name] in names:
                lines += [f"Notes : {line}" for line in _paragraphs(archive.read(notes_of[name]))]
            sections.append("\n".join(lines))
    return "\n\n".join(sections)


def _plain(payload: bytes, limits: ExtractionLimits) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    return payload.decode("latin-1")


_READERS = {
    "application/pdf": _pdf,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": _docx,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": _xlsx,
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": _pptx,
    "text/plain": _plain,
    "text/csv": _plain,
    "application/csv": _plain,
    "text/comma-separated-values": _plain,
}


def _tidy(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def attachment_text(content_type: str, payload: bytes | None,
                    limits: ExtractionLimits | None = None) -> tuple[str | None, str]:
    """The file's text (cut at the characters limit) and what happened, as a short French status.
    ``limits``: those of the process (its environment) unless given."""
    reader = _READERS.get(content_type)
    if reader is None or payload is None:
        return None, "non extrait"
    limits = limits or ExtractionLimits.from_environment()
    try:
        text = _tidy(reader(payload, limits))
    except InflatedTooLarge:
        return None, "trop volumineux une fois décompressé"
    except ImportError:
        return None, "lecteur indisponible"
    except Exception:  # noqa: BLE001 - a damaged file must not stop the archive
        return None, "illisible"
    if not text:
        return None, "sans texte (image ou scan ?)"
    if len(text) > limits.attachment_chars:
        return text[:limits.attachment_chars].rstrip() + "…", f"tronqué à {limits.attachment_chars} caractères"
    return text, "extrait"


def is_extractable(content_type: str) -> bool:
    return content_type in _READERS


__all__ = ["attachment_text", "is_extractable", "CONTENT_CHARS"]
