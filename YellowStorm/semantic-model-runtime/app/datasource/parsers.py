"""Deterministic bounded tabular discovery (Phase 3, P3.7-P3.10).

Pure functions over explicit bytes. CSV uses the stdlib; XLSX uses openpyxl
imported lazily inside the function body so importing the API never
initializes workbook parsing (see ``test_api_skeleton``).

Mirrors the NestJS ``SpreadsheetConceptResolver`` bounds where they exist:
50 MB pre-parse gate, 5000-row scan cap, 200-row profiles, ``Column N``
fallbacks, and ``__sheetRow`` row provenance.
"""

from __future__ import annotations

import csv
import hashlib
import io
import math
import zipfile
from datetime import date, datetime
from typing import Any

from .discovery import MAX_COMPRESSED_BYTES, validate_archive_safety

SHEET_ROW_KEY = "__sheetRow"

PREVIEW_SAMPLE_ROWS = 20
PREVIEW_MAX_COLUMNS = 50
PROFILE_ROW_LIMIT = 200
SCAN_ROW_LIMIT = 5000
SAMPLE_CELL_CHAR_LIMIT = 500
# Single-field bound: far above the stdlib default (128 KiB) so large valid
# cells parse and are shortened at sample time instead of failing the file.
MAX_CSV_FIELD_BYTES = 10 * 1024 * 1024

csv.field_size_limit(MAX_CSV_FIELD_BYTES)

_CSV_DELIMITERS = (",", ";", "\t", "|")


def _warn(warnings: list[dict[str, str]], code: str, message: str) -> None:
    warnings.append({"code": code, "message": message})


def _disambiguate_headers(raw: list[str], warnings: list[dict[str, str]]) -> list[str]:
    """Every final header is globally unique, and every rename warns.

    A rename is any deviation from the raw header text: empty or whitespace
    fallbacks, normalization, and ``__N`` collision suffixes alike. The
    ``__sheetRow`` provenance key is reserved up front; a source column with
    that name becomes the visible ``sheetRow`` (suffixed on further
    collisions) so its values survive in samples and field profiles, which
    exclude ``__*`` names by design.
    """
    used: set[str] = {SHEET_ROW_KEY}
    headers: list[str] = []
    renamed = False
    for index, name in enumerate(raw, start=1):
        base = name.strip() or f"Column {index}"
        if base == SHEET_ROW_KEY:
            base = "sheetRow"
        candidate, suffix = base, 2
        while candidate in used:
            candidate = f"{base}__{suffix}"
            suffix += 1
        if candidate != name:
            renamed = True
        used.add(candidate)
        headers.append(candidate)
    if renamed:
        _warn(warnings, "headers_renamed",
              "Column names were empty, normalized, or not unique; renamed copies carry a __N suffix.")
    return headers


def _truncate_cell(value: Any, warnings: list[dict[str, str]], state: dict[str, bool]) -> Any:
    if isinstance(value, str) and len(value) > SAMPLE_CELL_CHAR_LIMIT:
        if not state.get("truncated"):
            state["truncated"] = True
            _warn(warnings, "sample_cells_truncated",
                  f"Sample cell text is shortened to {SAMPLE_CELL_CHAR_LIMIT} characters.")
        return value[:SAMPLE_CELL_CHAR_LIMIT] + "…"
    return value


def _sample_cell(value: Any, warnings: list[dict[str, str]], state: dict[str, bool]) -> Any:
    """Serialize a typed record value for samples; profiles see the raw form."""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    return _truncate_cell(value, warnings, state)


def _infer_type(value: Any) -> str:
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return "number"
    if isinstance(value, (datetime, date)):
        return "date"
    return "text"


def field_profiles(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Mirror of the NestJS ``computeFieldProfiles`` over bounded rows."""
    names: list[str] = []
    for row in rows:
        for name in row:
            if not name.startswith("__") and name not in names:
                names.append(name)
    profiles = []
    for name in names:
        values = [row.get(name) for row in rows]
        populated = [v for v in values
                     if v is not None and not (isinstance(v, float) and math.isnan(v))
                     and Stringify(v) != ""]
        unique = {Stringify(v).strip().lower() for v in populated}
        first = populated[0] if populated else None
        profile: dict[str, Any] = {
            "name": name,
            "type": _infer_type(first),
            "sample": "" if first is None else Stringify(first)[:80],
            "populatedRatio": len(populated) / len(rows) if rows else 0,
            "uniqueRatio": len(unique) / len(populated) if populated else 0,
        }
        if profile["type"] == "text" and populated and all(_looks_numeric(v) for v in populated):
            profile["suggestedType"] = "number"
        profiles.append(profile)
    return profiles


def Stringify(value: Any) -> str:
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return ""
    return str(value)


def _looks_numeric(value: Any) -> bool:
    if isinstance(value, (int, float)):
        return True
    if not isinstance(value, str) or not value.strip():
        return False
    try:
        float(value.strip().replace(" ", "").replace(",", ".") if value.strip().count(",") == 1 and "." not in value else value.strip())
        return True
    except ValueError:
        return False


def _decode_text(data: bytes, options: dict[str, Any],
                 warnings: list[dict[str, str]]) -> tuple[str, str]:
    """Decode with explicit encoding or BOM-aware UTF-8, else latin-1."""
    override = options.get("encoding")
    if isinstance(override, str) and override:
        try:
            return data.decode(override), override
        except (LookupError, UnicodeDecodeError) as exc:
            raise ValueError("unsupported_encoding") from exc
    try:
        return data.decode("utf-8-sig"), "utf-8-sig"
    except UnicodeDecodeError:
        _warn(warnings, "encoding_fallback_latin1",
              "Content is not valid UTF-8; it was read as latin-1 and values must be confirmed.")
        return data.decode("latin-1"), "latin-1"


def _sniff_delimiter(text: str, options: dict[str, Any],
                     warnings: list[dict[str, str]]) -> str:
    """Quote-aware delimiter selection over a bounded prefix.

    Only the first 64 KiB are inspected, so a newline-dense file cannot force
    full-text materialization. Each candidate parses the first five non-empty
    lines with the csv reader (quotes honored), so a comma inside ``"..."``
    never votes for comma. The winner needs consistent multi-field rows, even
    for a header-only file; ties resolve in fixed ``, ; tab |`` order.
    """
    override = options.get("delimiter")
    if isinstance(override, str) and override in _CSV_DELIMITERS:
        return override
    prefix = text[:65536]
    sample = "\n".join([line for line in prefix.splitlines() if line.strip()][:5])
    best, best_fields = ",", 0
    for mark in _CSV_DELIMITERS:
        try:
            counts = [len(row) for row in csv.reader(io.StringIO(sample), delimiter=mark)]
        except csv.Error:
            continue
        if counts and len(set(counts)) == 1 and counts[0] > 1 and counts[0] > best_fields:
            best, best_fields = mark, counts[0]
    _warn(warnings, "delimiter_sniffed",
          f"Delimiter {best!r} was inferred from the first rows; confirm it before mapping.")
    return best


def _finalize(headers: list[str], records: list[dict[str, Any]], warnings: list[dict[str, str]],
              scanned: int, complete: bool, profile_rows: list[dict[str, Any]],
              structure: dict[str, Any]) -> dict[str, Any]:
    if not complete:
        _warn(warnings, "scan_capped",
              f"Reading stopped after {SCAN_ROW_LIMIT} rows; counts are estimates and the profile is partial.")
    samples: list[dict[str, Any]] = []
    trunc_state: dict[str, bool] = {}
    for record in records[:PREVIEW_SAMPLE_ROWS]:
        samples.append({key: _sample_cell(value, warnings, trunc_state)
                        for key, value in record.items()})
    return {
        "structure": structure,
        "samples": samples,
        "fieldProfiles": field_profiles(profile_rows[:PROFILE_ROW_LIMIT]),
        "warnings": warnings,
        "coverage": {"sampled": True, "completeProfileDone": complete},
        "scannedRows": scanned,
    }


def parse_csv_preview(data: bytes, options: dict[str, Any] | None = None) -> dict[str, Any]:
    """Bounded CSV preview (P3.10). Never guesses silently: encoding and
    delimiter decisions are recorded, and row 1 is an explicit header
    assumption, matching the existing NestJS mapping behavior."""
    opts = options or {}
    warnings: list[dict[str, str]] = []
    if len(data) > MAX_COMPRESSED_BYTES:
        raise ValueError("source_too_large")
    fingerprint = f"sha256:{hashlib.sha256(data).hexdigest()}"
    text, encoding = _decode_text(data, opts, warnings)
    delimiter = _sniff_delimiter(text, opts, warnings)

    records: list[dict[str, Any]] = []
    headers: list[str] = []
    scanned = 0
    complete = False
    header_row_number = 0
    try:
        reader = csv.reader(io.StringIO(text), delimiter=delimiter, strict=True)
        for record in reader:
            # Physical lines consumed, not records yielded: a quoted record
            # spanning thousands of lines must trip the bound the moment it
            # is consumed, or multiline input bypasses traversal limits.
            line_number = reader.line_num
            if any(cell.strip() for cell in record):
                if not headers:
                    header_row_number = line_number
                    headers = _disambiguate_headers(record, warnings)
                    _warn(warnings, "header_row_assumed",
                          f"Row {header_row_number} is assumed to hold headers; confirm before mapping.")
                    if len(headers) > PREVIEW_MAX_COLUMNS:
                        _warn(warnings, "columns_truncated",
                              f"Only the first {PREVIEW_MAX_COLUMNS} columns are previewed.")
                        headers = headers[:PREVIEW_MAX_COLUMNS]
                elif scanned < SCAN_ROW_LIMIT:
                    scanned += 1
                    row = {SHEET_ROW_KEY: line_number}
                    for index, name in enumerate(headers):
                        row[name] = record[index] if index < len(record) else ""
                    records.append(row)
            # Caps are checked at the end of the body so a reached limit
            # stops before the next record is requested: a multiline,
            # oversized, or malformed record starting beyond the bound is
            # never consumed. The bound is inclusive: a record ending exactly
            # on it consumed the whole budget, so traversal stops as partial
            # rather than spending another fetch to prove completeness.
            # A record already crossing the physical bound is processed once
            # (its content is field-limit bounded) and then stops.
            if scanned >= SCAN_ROW_LIMIT or line_number >= SCAN_ROW_LIMIT + 1:
                break
        else:
            complete = True
    except csv.Error as exc:
        if "field larger than field limit" in str(exc):
            raise ValueError("field_too_large") from exc
        raise ValueError("csv_malformed") from exc
    if not headers:
        raise ValueError("csv_no_header_row")

    structure = {"kind": "csv", "delimiter": delimiter, "encoding": encoding,
                 "bomPresent": data[:3] == b"\xef\xbb\xbf",
                 "headerRow": header_row_number, "columns": headers,
                 "dataRows": scanned if complete else None}
    return {"contentFingerprint": fingerprint,
            **_finalize(headers, records, warnings, scanned, complete, records, structure)}


def _xlsx_safety(data: bytes) -> list[str]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = archive.namelist()
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise ValueError("unreadable_source") from exc
    return [w["code"] for w in validate_archive_safety(
        detected_format="xlsx", compressed_bytes=len(data),
        decompressed_bytes=total, entry_count=len(names), paths=names)]


def _xlsx_value(value: Any) -> Any:
    """Keep typed values; samples serialize dates and drop non-finite floats."""
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return None if isinstance(value, float) and (math.isnan(value) or math.isinf(value)) else value
    if isinstance(value, (datetime, date)):
        return value
    return value if isinstance(value, str) else str(value)


def parse_xlsx_preview(data: bytes, options: dict[str, Any] | None = None) -> dict[str, Any]:
    """Bounded XLSX preview (P3.7-P3.9). Sheet dimensions are reported
    metadata until verified (W7); row 1 is an explicit header assumption."""
    import openpyxl

    opts = options or {}
    warnings: list[dict[str, str]] = []
    if len(data) > MAX_COMPRESSED_BYTES:
        raise ValueError("source_too_large")
    safety_codes = _xlsx_safety(data)
    if "source_too_large" in safety_codes or "archive_too_many_entries" in safety_codes:
        raise ValueError("source_too_large")
    _warn(warnings, "no_formula_execution",
          "Macros, formulas and external links are never executed.")
    fingerprint = f"sha256:{hashlib.sha256(data).hexdigest()}"

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("unreadable_source") from exc
    try:
        names = workbook.sheetnames
        if not names:
            raise ValueError("xlsx_no_sheets")
        requested = opts.get("sheetName")
        selected = requested if isinstance(requested, str) and requested in names else names[0]
        if isinstance(requested, str) and requested not in names:
            raise ValueError("sheet_not_found")
        sheets = [{"name": name,
                   "visible": workbook[name].sheet_state != "hidden",
                   # Reported only: read-only dimensions can be stale (W7).
                   "reportedRows": workbook[name].max_row,
                   "reportedColumns": workbook[name].max_column} for name in names]
        sheet = workbook[selected]

        header_cells = list(next(sheet.iter_rows(min_row=1, max_row=1,
                                                 values_only=True), []))
        # Width comes from populated header cells, not the worksheet-wide
        # max_column: trailing padding from wider data rows or stale
        # dimensions would otherwise invent synthetic Column N fields. One
        # row is cheap even at full Excel width.
        if any(v is not None and Stringify(v).strip() for v in header_cells):
            while header_cells and (header_cells[-1] is None
                                    or Stringify(header_cells[-1]).strip() == ""):
                header_cells.pop()
            raw_headers = ["" if v is None else Stringify(v) for v in header_cells]
        elif header_cells:
            # Header row exists but is entirely empty: positional fallback so
            # data below is still previewed instead of silently dropped.
            if len(header_cells) > PREVIEW_MAX_COLUMNS:
                _warn(warnings, "columns_truncated",
                      f"Only the first {PREVIEW_MAX_COLUMNS} columns are previewed.")
            raw_headers = [""] * min(len(header_cells), PREVIEW_MAX_COLUMNS)
        else:
            raw_headers = []
            _warn(warnings, "empty_sheet", "The selected sheet has no header row.")
        headers = _disambiguate_headers(raw_headers, warnings)
        _warn(warnings, "header_row_assumed",
              "Row 1 is assumed to hold headers; confirm before mapping.")
        width = len(headers)
        if width > PREVIEW_MAX_COLUMNS:
            _warn(warnings, "columns_truncated",
                  f"Only the first {PREVIEW_MAX_COLUMNS} columns are previewed.")
            headers = headers[:PREVIEW_MAX_COLUMNS]
            width = PREVIEW_MAX_COLUMNS

        records: list[dict[str, Any]] = []
        scanned = 0
        physical = 0
        complete = False
        for number, values in enumerate(
                sheet.iter_rows(min_row=2, max_col=max(width, 1),
                                values_only=True), start=2):
            # Physical traversal is bounded independently of non-empty rows and
            # of the reported dimensions, so a sparse sheet with a huge range
            # cannot force a full-workbook walk.
            physical += 1
            cells = list(values)
            if any(v is not None and Stringify(v).strip() for v in cells):
                if scanned < SCAN_ROW_LIMIT:
                    scanned += 1
                    row = {SHEET_ROW_KEY: number}
                    for index, name in enumerate(headers):
                        raw = cells[index] if index < len(cells) else None
                        row[name] = _xlsx_value(raw)
                    records.append(row)
            # Caps are checked at the end of the body so a reached limit stops
            # before the next worksheet row is requested. The physical bound
            # has no header allowance: this pass starts at row 2, unlike the
            # CSV line numbers. A file ending exactly on the budget therefore
            # reports partial rather than spending another row parse to prove
            # completeness; the batch preparation path remains the source of
            # exact counts.
            if scanned >= SCAN_ROW_LIMIT or physical >= SCAN_ROW_LIMIT:
                break
        else:
            complete = True

        formula_cells = _count_formula_cells(data, selected, max(width, 1))
        if formula_cells:
            _warn(warnings, "formula_values_are_cached",
                  f"{formula_cells} sampled cell(s) hold formulas; shown values are the cached "
                  "results, which may be stale or missing.")
    finally:
        workbook.close()

    structure = {"kind": "xlsx", "sheets": sheets, "selectedSheet": selected,
                 "headerRow": 1, "columns": headers,
                 "dataRows": scanned if complete else None}
    return {"contentFingerprint": fingerprint,
            **_finalize(headers, records, warnings, scanned, complete, records, structure)}


def _count_formula_cells(data: bytes, sheet_name: str, max_col: int) -> int:
    """Bounded formula check over the profiled prefix (P3.9)."""
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception:
        return 0
    try:
        sheet = workbook[sheet_name]
        count = 0
        for values in sheet.iter_rows(min_row=1, max_row=PROFILE_ROW_LIMIT + 1,
                                      max_col=max_col, values_only=False):
            for cell in values:
                if getattr(cell, "data_type", None) == "f":
                    count += 1
        return count
    except Exception:
        return 0
    finally:
        workbook.close()
